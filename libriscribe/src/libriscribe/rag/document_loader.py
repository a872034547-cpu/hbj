# src/libriscribe/rag/document_loader.py
"""文档解析器

支持 PDF/Word/Excel/TXT/Markdown 文件的解析和分块。
"""

import base64
import json
import logging
import hashlib
import os
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)


@dataclass
class DocumentChunk:
    """文档分块"""
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    chunk_id: str = ""

    def __post_init__(self):
        if not self.chunk_id:
            # 基于内容生成唯一 ID
            self.chunk_id = hashlib.md5(self.content.encode()).hexdigest()[:12]


class DocumentLoader:
    """文档加载和分块器"""

    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def load_file(self, file_path: str) -> List[DocumentChunk]:
        """加载单个文件并分块
        
        Args:
            file_path: 文件路径
            
        Returns:
            文档分块列表
        """
        path = Path(file_path)
        file_hash = self._file_hash(path)
        if not path.exists():
            logger.error(f"File not found: {file_path}")
            return []

        suffix = path.suffix.lower()
        text = ""

        try:
            if suffix == ".pdf":
                text = self._load_pdf(file_path)
            elif suffix == ".docx":
                text = self._load_docx(file_path)
            elif suffix == ".doc":
                text = self._load_doc(file_path)
            elif suffix in (".xlsx", ".xls"):
                text = self._load_excel(file_path)
            elif suffix in (".txt", ".md", ".markdown"):
                text = self._load_text(file_path)
            elif suffix in (".html", ".htm"):
                text = self._load_html(file_path)
            else:
                logger.warning(f"Unsupported file type: {suffix}")
                return []
        except Exception as e:
            logger.error(f"Error loading file {file_path}: {e}")
            return []

        text = self._clean_extracted_text(text)
        if not text.strip():
            logger.warning(f"No usable text extracted from {file_path}")
            return []

        # 分块
        chunks = self._split_text(text)

        # 添加可追溯元数据：source_id/evidence_id/chunk_hash 是后续引用核验的基础。
        source_id = hashlib.md5(f"{path.name}:{file_hash}".encode("utf-8", errors="ignore")).hexdigest()[:16]
        for i, chunk in enumerate(chunks):
            chunk_hash = hashlib.md5(chunk.content.encode("utf-8", errors="ignore")).hexdigest()
            evidence_id = hashlib.md5(f"{source_id}:{i}:{chunk_hash}".encode("utf-8", errors="ignore")).hexdigest()[:16]
            chunk.chunk_id = evidence_id
            chunk.metadata.update({
                "source": path.name,
                "source_id": source_id,
                "evidence_id": evidence_id,
                "file_path": str(path.absolute()),
                "file_type": suffix,
                "file_hash": file_hash,
                "chunk_hash": chunk_hash,
                "chunk_index": i,
                "total_chunks": len(chunks),
                "char_count": len(chunk.content),
            })

        logger.info(f"Loaded {len(chunks)} chunks from {file_path}")
        return chunks

    def load_directory(self, dir_path: str, recursive: bool = True) -> List[DocumentChunk]:
        """加载目录下的所有支持文件
        
        Args:
            dir_path: 目录路径
            recursive: 是否递归子目录
            
        Returns:
            所有文件的文档分块列表
        """
        path = Path(dir_path)
        if not path.is_dir():
            logger.error(f"Directory not found: {dir_path}")
            return []

        supported_extensions = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt", ".md", ".markdown", ".html", ".htm"}
        all_chunks = []

        pattern = "**/*" if recursive else "*"
        for file_path in sorted(path.glob(pattern)):
            if file_path.is_file() and file_path.suffix.lower() in supported_extensions:
                chunks = self.load_file(str(file_path))
                all_chunks.extend(chunks)

        logger.info(f"Loaded {len(all_chunks)} total chunks from {dir_path}")
        return all_chunks

    def _load_pdf(self, file_path: str) -> str:
        """解析 PDF 文件。

        优先走本地文本抽取；如果 PDF 是扫描件、unstructured 解析异常，或 PyPDF2 抽取为空，
        再按需调用 PaddleOCR 在线接口做最小化文字识别兜底。
        """
        text = ""

        try:
            from unstructured.partition.pdf import partition_pdf
            elements = partition_pdf(filename=file_path)
            text = "\n\n".join([str(el) for el in elements if str(el).strip()])
            if self._clean_extracted_text(text).strip():
                return text
            logger.warning("unstructured PDF parsing produced no usable text for %s; trying PyPDF2 fallback", file_path)
        except ImportError:
            logger.warning("unstructured not installed, trying PyPDF2 fallback")
        except Exception as e:
            logger.warning("unstructured PDF parsing failed for %s: %s; trying PyPDF2 fallback", file_path, e)

        try:
            import PyPDF2
            parts = []
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        parts.append(page_text)
            text = "\n\n".join(parts)
            if self._clean_extracted_text(text).strip():
                return text
            logger.warning("PyPDF2 produced no usable text for %s; trying PaddleOCR fallback", file_path)
        except ImportError:
            logger.warning("PyPDF2 not installed; trying PaddleOCR fallback")
        except Exception as e:
            logger.warning("PyPDF2 PDF parsing failed for %s: %s; trying PaddleOCR fallback", file_path, e)

        ocr_text = self._load_pdf_with_paddleocr(file_path)
        if ocr_text.strip():
            return ocr_text
        logger.error("No PDF parser succeeded for %s. Install unstructured/PyPDF2 or configure PaddleOCR token.", file_path)
        return ""

    def _load_pdf_with_paddleocr(self, file_path: str) -> str:
        """使用 PaddleOCR AIStudio layout-parsing 接口识别 PDF 文字，返回 Markdown/纯文本。

        配置方式：设置环境变量 PADDLEOCR_AISTUDIO_TOKEN；可选设置
        PADDLEOCR_AISTUDIO_LAYOUT_URL、PADDLEOCR_AISTUDIO_TIMEOUT_SECONDS。
        AIStudio 官方示例要求使用 JSON + base64 文件内容，并以 ``Authorization: token ...`` 鉴权。
        """
        token = os.getenv("PADDLEOCR_AISTUDIO_TOKEN", "").strip()
        if not token:
            logger.warning("PaddleOCR fallback skipped: PADDLEOCR_AISTUDIO_TOKEN is not configured")
            return ""

        try:
            import requests
        except ImportError:
            logger.warning("PaddleOCR fallback skipped: requests is not installed")
            return ""

        api_url = (
            os.getenv("PADDLEOCR_AISTUDIO_LAYOUT_URL", "")
            or os.getenv("PADDLEOCR_AISTUDIO_JOB_URL", "")
            or "https://c3q5maa8cfq9rex9.aistudio-app.com/layout-parsing"
        ).strip()
        timeout_seconds = max(30, self._safe_int(os.getenv("PADDLEOCR_AISTUDIO_TIMEOUT_SECONDS"), 180))
        headers = {
            "Authorization": f"token {token}",
            "Content-Type": "application/json",
        }

        try:
            with open(file_path, "rb") as f:
                file_data = base64.b64encode(f.read()).decode("ascii")

            response = requests.post(
                api_url,
                json={
                    "file": file_data,
                    "fileType": 0,
                    "useDocOrientationClassify": False,
                    "useDocUnwarping": False,
                    "useChartRecognition": False,
                },
                headers=headers,
                timeout=timeout_seconds,
            )
            if response.status_code != 200:
                logger.warning("PaddleOCR layout parsing failed for %s: %s %s", file_path, response.status_code, response.text[:500])
                return ""

            payload = response.json()
            result = payload.get("result") or {}
            pages: list[str] = []
            for item in result.get("layoutParsingResults") or []:
                markdown = item.get("markdown") or {}
                page_text = str(markdown.get("text") or "").strip()
                if page_text:
                    pages.append(page_text)

            if pages:
                return "\n\n".join(self._dedupe_text_lines(pages))

            fallback_text = self._collect_text_from_ocr_object(result)
            if fallback_text.strip():
                return fallback_text

            logger.warning("PaddleOCR layout parsing returned no text for %s: %s", file_path, response.text[:500])
            return ""
        except Exception as e:
            logger.warning("PaddleOCR fallback failed for %s: %s", file_path, e)
            return ""

    def _extract_paddleocr_jsonl_text(self, jsonl_text: str) -> str:
        """从 PaddleOCR JSONL 结果中尽量抽取纯文本，不下载图片。"""
        pages: list[str] = []
        for raw_line in str(jsonl_text or "").splitlines():
            line = raw_line.strip()
            if not line:
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                continue
            result = record.get("result", record)
            page_text = self._collect_text_from_ocr_object(result)
            if page_text.strip():
                pages.append(page_text.strip())
        return "\n\n".join(self._dedupe_text_lines(pages))

    def _collect_text_from_ocr_object(self, value: Any) -> str:
        """兼容不同 PaddleOCR 返回结构，递归提取 recText/text/content 等文字字段。"""
        text_fields = {"recText", "text", "content", "transcription", "label", "words"}
        parts: list[str] = []
        if isinstance(value, dict):
            for key in ("recTexts", "texts", "textLines", "ocrResults", "prunedResult"):
                nested = value.get(key)
                if nested is not None:
                    nested_text = self._collect_text_from_ocr_object(nested)
                    if nested_text:
                        parts.append(nested_text)
            for key, nested in value.items():
                if key in text_fields and isinstance(nested, str) and nested.strip():
                    parts.append(nested.strip())
                elif key not in {"ocrImage", "image", "bbox", "points", "score"} and isinstance(nested, (dict, list, tuple)):
                    nested_text = self._collect_text_from_ocr_object(nested)
                    if nested_text:
                        parts.append(nested_text)
        elif isinstance(value, (list, tuple)):
            for item in value:
                if isinstance(item, str) and item.strip():
                    parts.append(item.strip())
                    continue
                nested_text = self._collect_text_from_ocr_object(item)
                if nested_text:
                    parts.append(nested_text)
        elif isinstance(value, str) and value.strip():
            parts.append(value.strip())
        return "\n".join(self._dedupe_text_lines(parts)).strip()

    @staticmethod
    def _dedupe_text_lines(lines: list[str]) -> list[str]:
        """保持顺序去重 OCR 文本，避免递归兼容解析造成重复片段。"""
        deduped: list[str] = []
        seen: set[str] = set()
        for item in lines:
            for raw_line in str(item or "").splitlines():
                line = raw_line.strip()
                if not line or line in seen:
                    continue
                seen.add(line)
                deduped.append(line)
        return deduped

    @staticmethod
    def _safe_int(value: Any, default: int) -> int:
        try:
            return int(value)
        except (TypeError, ValueError):
            return default

    def _load_docx(self, file_path: str) -> str:
        """解析 Word 文件。

        优先按真实 DOCX ZIP 包解析；如果用户上传的是误命名为 .docx 的纯文本、
        Markdown 或其他可解码文本，则自动按文本读取，避免资料库索引直接失败。
        """
        try:
            from unstructured.partition.docx import partition_docx
            elements = partition_docx(filename=file_path)
            text = "\n\n".join([str(el) for el in elements if str(el).strip()])
            if text.strip():
                return text
        except ImportError:
            pass
        except Exception as e:
            logger.warning(f"unstructured DOCX parsing failed for {file_path}: {e}; trying python-docx/text fallback")

        try:
            from docx import Document
            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n\n".join(paragraphs)
            if text.strip():
                return text
        except ImportError:
            logger.warning("python-docx not installed; trying text fallback")
        except Exception as e:
            logger.warning(f"python-docx parsing failed for {file_path}: {e}; trying text fallback")

        fallback_text = self._load_text_fallback(file_path)
        if fallback_text.strip():
            logger.warning(f"Loaded {file_path} as plain text fallback because it is not a valid DOCX package")
            return fallback_text

        logger.error("No DOCX parser succeeded and text fallback produced no content")
        return ""

    def _load_doc(self, file_path: str) -> str:
        """解析旧版 Word .doc（OLE/CFB）文件。

        .doc 是二进制复合文档，不能按普通文本兜底，否则会把 OLE 结构索引成乱码。
        当前环境没有稳定 .doc 解析依赖时，直接返回空并提示用户转换格式。
        """
        logger.warning(
            f"Legacy .doc/OLE file is not indexed directly: {file_path}. "
            "Please save it as .docx, .pdf or .txt before uploading."
        )
        return ""

    def _load_excel(self, file_path: str) -> str:
        """解析 Excel 文件"""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True)
            text_parts = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                text_parts.append(f"Sheet: {sheet}")
                for row in ws.iter_rows(values_only=True):
                    row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                    if row_text.strip(" | "):
                        text_parts.append(row_text)
            return "\n".join(text_parts)
        except ImportError:
            logger.error("openpyxl not installed for Excel parsing")
            return ""

    def _load_text(self, file_path: str) -> str:
        """解析纯文本/Markdown 文件"""
        return self._load_text_fallback(file_path)

    def _load_text_fallback(self, file_path: str) -> str:
        """尽量以常见文本编码读取文件，用于误命名文档兜底。"""
        for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk", "latin-1"):
            try:
                with open(file_path, "r", encoding=encoding, errors="strict") as f:
                    text = f.read()
                if text.strip():
                    return text
            except UnicodeDecodeError:
                continue
            except Exception as e:
                logger.warning(f"Text fallback failed for {file_path} with {encoding}: {e}")
                break
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        except Exception as e:
            logger.warning(f"Unable to read text fallback for {file_path}: {e}")
            return ""

    def _load_html(self, file_path: str) -> str:
        """解析 HTML 文件"""
        try:
            from bs4 import BeautifulSoup
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                soup = BeautifulSoup(f.read(), "html.parser")
                return soup.get_text(separator="\n", strip=True)
        except ImportError:
            logger.error("beautifulsoup4 not installed for HTML parsing")
            return ""

    def _is_probably_binary_text(self, text: str) -> bool:
        """判断兜底读取结果是否明显是二进制乱码。"""
        if not text:
            return True
        sample = text[:4000]
        if sample.startswith("ÐÏ\x11à¡±\x1aá") or "R\x00o\x00o\x00t\x00 \x00E\x00n\x00t\x00r\x00y" in sample:
            return True
        if "\x00" in sample:
            return True
        replacement_ratio = sample.count("�") / max(len(sample), 1)
        if replacement_ratio > 0.05:
            return True
        readable = re.findall(r"[\u4e00-\u9fffA-Za-z0-9，。；：、！？（）《》\[\]【】\s\-_/,.:%]", sample)
        readable_ratio = len(readable) / max(len(sample), 1)
        return readable_ratio < 0.35

    def _clean_extracted_text(self, text: str) -> str:
        """清理抽取文本并过滤明显乱码。"""
        if self._is_probably_binary_text(text):
            return ""
        text = text.replace("\x00", "")
        text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", "", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        if self._looks_like_office_binary_garbage(text):
            return ""
        return text.strip()

    def _looks_like_office_binary_garbage(self, text: str) -> bool:
        """识别被错误解码后的 Office/OLE/WPS 二进制结构文本。"""
        if not text:
            return True
        sample = text[:8000]
        compact = re.sub(r"[\x00\s�\ufffd]+", "", sample).lower()
        office_markers = (
            "rootentry",
            "worddocument",
            "normal.dotm",
            "wpsoffice",
            "administrator",
            "ole",
            "compobj",
        )
        marker_hits = sum(1 for marker in office_markers if marker in compact)
        if marker_hits >= 2:
            return True
        if sample.count("þÿÿ") >= 3 or sample.count("���") >= 20:
            return True
        # 旧版 Word/WPS 二进制误读常出现大量孤立控制字符与低信息量符号。
        symbol_count = len(re.findall(r"[�\ufffdþÿÐÏà¡±]", sample))
        if symbol_count / max(len(sample), 1) > 0.08 and marker_hits >= 1:
            return True
        return False

    def _file_hash(self, path: Path) -> str:
        """计算文件 hash，用于资料来源去重和证据追溯。"""
        h = hashlib.md5()
        try:
            with open(path, "rb") as f:
                for block in iter(lambda: f.read(1024 * 1024), b""):
                    h.update(block)
            return h.hexdigest()
        except Exception:
            return hashlib.md5(str(path).encode("utf-8", errors="ignore")).hexdigest()

    def is_usable_text(self, text: str) -> bool:
        """供 Web 层/测试使用：判断文本是否适合进入证据库。"""
        return bool(self._clean_extracted_text(text).strip())

    def _split_text(self, text: str) -> List[DocumentChunk]:
        """将文本分块
        
        使用滑动窗口策略，按段落边界分割。
        """
        if len(text) <= self.chunk_size:
            return [DocumentChunk(content=text)]

        chunks = []
        # 先按段落分割
        paragraphs = text.split("\n\n")
        current_chunk = ""

        for para in paragraphs:
            para = para.strip()
            if not para:
                continue

            # 如果当前块加上新段落不超过限制
            if len(current_chunk) + len(para) + 2 <= self.chunk_size:
                current_chunk = current_chunk + "\n\n" + para if current_chunk else para
            else:
                # 保存当前块
                if current_chunk:
                    chunks.append(DocumentChunk(content=current_chunk.strip()))

                # 如果单个段落超过限制，按字符数强制分割
                if len(para) > self.chunk_size:
                    sub_chunks = self._force_split(para)
                    chunks.extend(sub_chunks)
                    current_chunk = ""
                else:
                    current_chunk = para

        # 处理最后一个块
        if current_chunk.strip():
            chunks.append(DocumentChunk(content=current_chunk.strip()))

        # 添加重叠
        if self.chunk_overlap > 0 and len(chunks) > 1:
            chunks = self._add_overlap(chunks)

        return chunks

    def _force_split(self, text: str) -> List[DocumentChunk]:
        """强制按字符数分割长文本"""
        chunks = []
        start = 0
        while start < len(text):
            end = start + self.chunk_size
            chunk_text = text[start:end]
            chunks.append(DocumentChunk(content=chunk_text.strip()))
            start = end - self.chunk_overlap
        return chunks

    def _add_overlap(self, chunks: List[DocumentChunk]) -> List[DocumentChunk]:
        """为分块添加重叠内容"""
        if len(chunks) <= 1:
            return chunks

        overlapped = [chunks[0]]
        for i in range(1, len(chunks)):
            prev_content = chunks[i - 1].content
            # 取前一个块的末尾作为重叠
            overlap_text = prev_content[-self.chunk_overlap:] if len(prev_content) > self.chunk_overlap else prev_content
            new_content = overlap_text + " " + chunks[i].content
            # 确保不超过限制
            if len(new_content) > self.chunk_size * 1.2:
                new_content = new_content[:self.chunk_size]
            overlapped.append(DocumentChunk(content=new_content.strip(), metadata=chunks[i].metadata.copy(), chunk_id=chunks[i].chunk_id))
        return overlapped