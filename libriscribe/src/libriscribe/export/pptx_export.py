"""PPTX 导出器。

将专著正文转换为适合汇报/路演的 PowerPoint：封面、目录、章节摘要、章节要点、结束页。
默认采用中文演示文档常用的 16:9、深绿/米白配色和清晰层级，避免把整篇长文直接塞进幻灯片。
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Iterable, List

from libriscribe.utils.chinese_labels import format_chapter_label

logger = logging.getLogger(__name__)


class PptxExporter:
    """导出项目为 PPTX 演示文稿。"""

    def __init__(self) -> None:
        try:
            import pptx  # noqa: F401
            self._pptx_available = True
        except ImportError:
            logger.error("python-pptx not installed")
            self._pptx_available = False

    def export(
        self,
        chapters: Iterable[dict],
        output_path: str,
        title: str = "",
        author: str = "",
        genre: str = "",
        language: str = "简体中文",
        reference_style: str = "academic",
    ) -> None:
        """导出 PPTX。

        reference_style 当前提供内置中文学术汇报版式；保留参数用于后续接入用户上传模板。
        """
        if not self._pptx_available:
            raise ImportError("python-pptx is required. Please install it with: pip install python-pptx")

        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        chapter_list = list(chapters)
        self._add_title_slide(prs, title or "未命名书稿", author, genre)
        self._add_toc_slide(prs, chapter_list)

        for chapter in chapter_list:
            chapter_title = self._chapter_title(chapter)
            content = self._strip_markdown(chapter.get("content", ""))
            summary = self._summarize_text(content, max_chars=220)
            bullets = self._extract_bullets(content, max_items=5)
            self._add_chapter_overview_slide(prs, chapter_title, summary, bullets)

            section_bullets = self._extract_section_bullets(chapter.get("content", ""), max_sections=4)
            if section_bullets:
                self._add_section_slide(prs, chapter_title, section_bullets)

        self._add_closing_slide(prs, title or "书稿汇报")
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path)
        logger.info("PPTX exported to: %s", output_path)

    def _add_title_slide(self, prs, title: str, author: str, genre: str) -> None:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._set_bg(slide, "F7F1E5")
        self._add_accent_bar(slide)
        box = slide.shapes.add_textbox(Inches(1.05), Inches(1.45), Inches(11.2), Inches(2.2))
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(36)
        run.font.bold = True
        run.font.color.rgb = self._rgb("243326")

        sub = slide.shapes.add_textbox(Inches(1.6), Inches(3.85), Inches(10.1), Inches(1.2))
        tf2 = sub.text_frame
        tf2.clear()
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = " · ".join([item for item in [genre, author] if item]) or "专著汇报"
        r2.font.name = "Microsoft YaHei"
        r2.font.size = Pt(18)
        r2.font.color.rgb = self._rgb("496D57")

    def _add_toc_slide(self, prs, chapters: List[dict]) -> None:
        from pptx.util import Inches, Pt

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._set_bg(slide, "FFFDF7")
        self._add_slide_title(slide, "目录 / 汇报结构")
        left = Inches(1.0)
        top = Inches(1.55)
        width = Inches(11.2)
        height = Inches(5.25)
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.clear()
        for idx, chapter in enumerate(chapters[:12], start=1):
            p = tf.paragraphs[0] if idx == 1 else tf.add_paragraph()
            p.text = f"{idx:02d}  {self._chapter_title(chapter)}"
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(18)
            p.font.color.rgb = self._rgb("243326")
            p.space_after = Pt(8)

    def _add_chapter_overview_slide(self, prs, title: str, summary: str, bullets: List[str]) -> None:
        from pptx.util import Inches, Pt

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._set_bg(slide, "FFFDF7")
        self._add_slide_title(slide, title[:42])

        summary_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.35), Inches(11.6), Inches(1.35))
        summary_tf = summary_box.text_frame
        summary_tf.clear()
        p = summary_tf.paragraphs[0]
        p.text = summary or "本章暂无可提炼摘要。"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(15)
        p.font.color.rgb = self._rgb("4E5C50")

        bullet_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.95), Inches(10.9), Inches(3.8))
        tf = bullet_box.text_frame
        tf.clear()
        for idx, bullet in enumerate(bullets or ["建议补充本章关键观点、资料证据和结论。"], start=1):
            p = tf.paragraphs[0] if idx == 1 else tf.add_paragraph()
            p.text = bullet[:90]
            p.level = 0
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(18)
            p.font.color.rgb = self._rgb("243326")
            p.space_after = Pt(10)

    def _add_section_slide(self, prs, chapter_title: str, section_bullets: List[str]) -> None:
        from pptx.util import Inches, Pt

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._set_bg(slide, "F7F1E5")
        self._add_slide_title(slide, f"{chapter_title[:30]}：章节要点")
        box = slide.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.2), Inches(5.3))
        tf = box.text_frame
        tf.clear()
        for idx, item in enumerate(section_bullets, start=1):
            p = tf.paragraphs[0] if idx == 1 else tf.add_paragraph()
            p.text = item[:105]
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(17)
            p.font.color.rgb = self._rgb("243326")
            p.space_after = Pt(12)

    def _add_closing_slide(self, prs, title: str) -> None:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._set_bg(slide, "496D57")
        box = slide.shapes.add_textbox(Inches(1.2), Inches(2.45), Inches(10.9), Inches(1.4))
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "谢谢 / Thank You"
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(34)
        r.font.bold = True
        r.font.color.rgb = self._rgb("FFFDF7")

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = title
        r2.font.name = "Microsoft YaHei"
        r2.font.size = Pt(16)
        r2.font.color.rgb = self._rgb("F7F1E5")

    def _add_slide_title(self, slide, title: str) -> None:
        from pptx.util import Inches, Pt

        self._add_accent_bar(slide)
        box = slide.shapes.add_textbox(Inches(0.75), Inches(0.38), Inches(11.9), Inches(0.65))
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = self._rgb("243326")

    def _add_accent_bar(self, slide) -> None:
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches

        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._rgb("496D57")
        shape.line.fill.background()

    def _set_bg(self, slide, color: str) -> None:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = self._rgb(color)

    def _rgb(self, color: str):
        from pptx.dml.color import RGBColor

        color = color.strip().lstrip("#")
        return RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))

    def _chapter_title(self, chapter: dict) -> str:
        number = chapter.get("number") or chapter.get("chapter_number") or ""
        title = str(chapter.get("display_title") or chapter.get("title") or "未命名章节").strip()
        if number and not title.startswith(("第", "Chapter")):
            return format_chapter_label(number, title)
        return title

    def _strip_markdown(self, text: str) -> str:
        text = re.sub(r"```.*?```", "", str(text or ""), flags=re.S)
        text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
        text = re.sub(r"[*_`>]+", "", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    def _summarize_text(self, text: str, max_chars: int = 220) -> str:
        clean = re.sub(r"\s+", " ", self._strip_markdown(text)).strip()
        return clean[:max_chars].rstrip() + ("…" if len(clean) > max_chars else "")

    def _extract_bullets(self, text: str, max_items: int = 5) -> List[str]:
        clean = self._strip_markdown(text)
        candidates = []
        for line in clean.splitlines():
            line = re.sub(r"^[-*\d.、\s]+", "", line).strip()
            if 18 <= len(line) <= 140 and "本章实际字数" not in line:
                candidates.append(line)
        if len(candidates) < max_items:
            parts = re.split(r"(?<=[。！？；])", re.sub(r"\s+", "", clean))
            candidates.extend([p.strip() for p in parts if 18 <= len(p.strip()) <= 140])
        seen = []
        for item in candidates:
            if item not in seen:
                seen.append(item)
            if len(seen) >= max_items:
                break
        return seen

    def _extract_section_bullets(self, markdown: str, max_sections: int = 4) -> List[str]:
        lines = []
        for line in str(markdown or "").splitlines():
            m = re.match(r"^#{2,4}\s+(.+)$", line.strip())
            if m:
                title = m.group(1).strip()
                if "本章小结" not in title and "参考文献" not in title:
                    lines.append(f"• {title}")
            if len(lines) >= max_sections:
                break
        return lines
