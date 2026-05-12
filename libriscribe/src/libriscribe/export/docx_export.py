# src/libriscribe/export/docx_export.py
"""DOCX 导出"""

import logging
import re
from pathlib import Path
from typing import List

from libriscribe.utils.chinese_labels import format_chapter_label, strip_leading_chapter_heading

logger = logging.getLogger(__name__)


class DocxExporter:
    def __init__(self):
        try:
            from docx import Document
            self._docx_available = True
        except ImportError:
            logger.error("python-docx not installed")
            self._docx_available = False

    def export(self, chapters: List[dict], output_path: str, title: str = "", author: str = "", genre: str = "", language: str = "English"):
        if not self._docx_available:
            raise ImportError("python-docx is required")
        from docx import Document

        doc = Document()
        self._apply_reference_layout(doc)

        self._add_title_page(doc, title, author, genre)
        doc.add_page_break()
        doc.add_heading('目录', level=1)
        doc.add_paragraph('（请在 Word 中更新目录：引用 -> 更新目录）')
        doc.add_page_break()

        for chapter in chapters:
            self._add_chapter(doc, chapter)

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        logger.info(f"DOCX exported to: {output_path}")

    def _apply_reference_layout(self, doc):
        """套用用户参考 Word 的基础版式：A4、左右 3.17cm、上下 2.54cm、中文专著标题层级。"""
        from docx.enum.section import WD_SECTION_START
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt

        section = doc.sections[0]
        section.start_type = WD_SECTION_START.NEW_PAGE
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

        def set_font(style_name: str, *, size: float, bold=None, align=None, first_indent: float | None = None, line_spacing: float | None = None, before: float = 0, after: float = 0):
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            style.font.size = Pt(size)
            style.font.bold = bold
            pf = style.paragraph_format
            pf.alignment = align
            pf.space_before = Pt(before)
            pf.space_after = Pt(after)
            if first_indent is not None:
                pf.first_line_indent = Pt(first_indent)
            if line_spacing is not None:
                pf.line_spacing = line_spacing
            return style

        set_font("Normal", size=10.5, bold=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=21, line_spacing=1.15, after=0)
        set_font("Title", size=36, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=156, after=0)
        set_font("Heading 1", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0, before=16.5, after=16.5)
        set_font("Heading 2", size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0, line_spacing=1.0, before=3, after=3)
        set_font("Heading 3", size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21, line_spacing=1.15, before=1.5, after=1.5)
        set_font("Heading 4", size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21, line_spacing=1.15, before=0, after=0)
        if "Body Text" in [s.name for s in doc.styles]:
            set_font("Body Text", size=10.5, bold=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=21, line_spacing=1.15, after=0)

    def _add_title_page(self, doc, title, author, genre):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        for _ in range(5):
            doc.add_paragraph('')
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title or "未命名书稿")
        r.bold = True
        self._set_run_cn_font(r, 36)
        if genre:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run(genre)
            self._set_run_cn_font(r2, 15)
        if author:
            p3 = doc.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r3 = p3.add_run(f'作者：{author}')
            self._set_run_cn_font(r3, 12)
        doc.add_page_break()

    def _add_chapter(self, doc, chapter):
        number = chapter.get('number') or chapter.get('chapter_number')
        title = chapter.get('title', '未命名章节')
        display_title = chapter.get('display_title') or title
        content = chapter.get('content', '')
        if number:
            if int(number) < 1:
                raise ValueError(f"Invalid chapter number for DOCX export: {number}")
            title = chapter.get('title', f'Chapter {number}')
            display_title = chapter.get('display_title') or format_chapter_label(number, title)
            content = content and self._strip_leading_chapter_heading(content, number, title) or ""
        else:
            content = self._strip_leading_part_heading(content, display_title)
        doc.add_heading(display_title, level=1)
        self._add_markdown_content(content or "", doc)
        doc.add_page_break()

    def _strip_leading_part_heading(self, content, title):
        pattern = rf"^\s*#{{1,6}}\s*{re.escape(str(title).strip())}\s*\n+"
        return re.sub(pattern, "", str(content or ""), count=1)

    def _strip_leading_chapter_heading(self, content, number, title):
        return strip_leading_chapter_heading(content, number, title)

    def _set_run_cn_font(self, run, size: float | None = None):
        from docx.oxml.ns import qn
        from docx.shared import Pt
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if size is not None:
            run.font.size = Pt(size)

    def _add_markdown_content(self, content, doc):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        lines = content.split('\n')
        in_code = False
        code_buf = []
        for line in lines:
            if line.strip().startswith('```'):
                if in_code:
                    p = doc.add_paragraph()
                    r = p.add_run('\n'.join(code_buf))
                    r.font.name = 'Courier New'
                    r.font.size = Pt(10)
                    code_buf = []
                    in_code = False
                else:
                    in_code = True
                continue
            if in_code:
                code_buf.append(line)
                continue
            if line.startswith('#### '):
                doc.add_heading(line[5:].strip(), level=4)
            elif line.startswith('### '):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith('## '):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:].strip(), level=1)
            elif line.strip().startswith('- ') or line.strip().startswith('* '):
                doc.add_paragraph(line.strip()[2:], style='List Bullet')
            elif not line.strip():
                continue
            else:
                style_name = 'Body Text' if 'Body Text' in [s.name for s in doc.styles] else None
                para = doc.add_paragraph(style=style_name)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.first_line_indent = Pt(21)
                para.paragraph_format.line_spacing = 1.15
                parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', line.strip())
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        r = para.add_run(part[2:-2])
                        r.bold = True
                    elif part.startswith('*') and part.endswith('*'):
                        r = para.add_run(part[1:-1])
                        r.italic = True
                    elif part.startswith('`') and part.endswith('`'):
                        r = para.add_run(part[1:-1])
                        r.font.name = 'Courier New'
                    else:
                        r = para.add_run(part)
                    self._set_run_cn_font(r, 10.5)

    def export_from_project(self, project_dir, output_path, title="", author="", genre="", language="English"):
        from libriscribe.utils.file_utils import get_chapter_files, read_markdown_file
        chapter_files = get_chapter_files(project_dir)
        chapters = []
        for i, cf in enumerate(chapter_files, 1):
            content = read_markdown_file(cf)
            m = re.search(r'^#+\s*(.+)$', content, re.MULTILINE)
            chapters.append({"number": i, "title": m.group(1) if m else f"Chapter {i}", "content": content})
        self.export(chapters, output_path, title, author, genre, language)
