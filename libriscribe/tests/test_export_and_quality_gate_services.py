from __future__ import annotations

import importlib.util
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from libriscribe.export.docx_export import DocxExporter
from libriscribe.knowledge_base import (
    Chapter,
    ChapterSection,
    EvidenceChunk,
    ProjectKnowledgeBase,
    SourceDocument,
    TaskLog,
    Citation,
)
from libriscribe.services.citation_service import CitationService
from libriscribe.services.export_service import ExportService
from libriscribe.services.pipeline_service import PipelineService
from libriscribe.services.quality_service import QualityService
from libriscribe.services.source_service import SourceService
from libriscribe.utils.academic_prompt import finalize_academic_chapter


def test_export_service_delegates_to_exporters_and_returns_paths(monkeypatch: Any, tmp_path: Path) -> None:
    calls: list[tuple[str, list[dict[str, Any]], str, dict[str, Any]]] = []

    def _write_stub(path: str, suffix: str) -> None:
        actual = Path(path if suffix != ".tex" or str(path).endswith(".tex") else f"{path}.tex")
        actual.parent.mkdir(parents=True, exist_ok=True)
        actual.write_bytes(("stub-export-content-" + suffix) .encode("utf-8") * 80)

    def fake_docx_export(self: Any, chapters: list[dict[str, Any]], output_path: str, **kwargs: Any) -> None:
        calls.append(("docx", chapters, output_path, kwargs))
        _write_stub(output_path, ".docx")

    def fake_pdf_export(self: Any, chapters: list[dict[str, Any]], output_path: str, **kwargs: Any) -> None:
        calls.append(("pdf", chapters, output_path, kwargs))
        _write_stub(output_path, ".pdf")

    def fake_pptx_export(self: Any, chapters: list[dict[str, Any]], output_path: str, **kwargs: Any) -> None:
        calls.append(("pptx", chapters, output_path, kwargs))
        _write_stub(output_path, ".pptx")

    def fake_latex_export(self: Any, chapters: list[dict[str, Any]], output_path: str, **kwargs: Any) -> None:
        calls.append(("latex", chapters, output_path, kwargs))
        _write_stub(output_path, ".tex")

    monkeypatch.setattr("libriscribe.export.docx_export.DocxExporter.export", fake_docx_export)
    monkeypatch.setattr("libriscribe.export.pdf_export.PdfExporter.export", fake_pdf_export)
    monkeypatch.setattr("libriscribe.export.pptx_export.PptxExporter.export", fake_pptx_export)
    monkeypatch.setattr("libriscribe.export.latex_export.LatexExporter.export", fake_latex_export)

    chapters = [{"chapter_number": 1, "title": "第一章", "content": "正文"}]
    service = ExportService()

    docx_path = tmp_path / "book.docx"
    pdf_path = tmp_path / "book.pdf"
    pptx_path = tmp_path / "book.pptx"
    latex_base_path = tmp_path / "book"

    assert service.export_docx(chapters, docx_path, title="书名", author="作者") == docx_path
    assert service.export_pdf(chapters, pdf_path, title="书名", genre="专著") == pdf_path
    assert service.export_pptx(chapters, pptx_path, title="书名", language="简体中文") == pptx_path
    assert service.export_latex(chapters, latex_base_path, title="书名", language="zh-CN") == Path(f"{latex_base_path}.tex")

    assert [call[0] for call in calls] == ["docx", "pdf", "pptx", "latex"]
    assert calls[0][1] == chapters
    assert calls[0][2] == str(docx_path)
    assert calls[1][2] == str(pdf_path)
    assert calls[2][2] == str(pptx_path)
    assert calls[3][2] == str(latex_base_path)
    assert calls[0][3]["title"] == "书名"
    assert calls[0][3]["author"] == "作者"
    assert calls[1][3]["genre"] == "专著"
    assert calls[2][3]["language"] == "简体中文"
    assert calls[3][3]["language"] == "zh-CN"


def test_export_service_validates_missing_or_tiny_outputs(tmp_path: Path) -> None:
    service = ExportService()
    missing = tmp_path / "missing.docx"
    tiny = tmp_path / "tiny.pdf"
    tiny.write_bytes(b"x")

    missing_report = service.validate_export_file(missing, expected_suffix=".docx")
    tiny_report = service.validate_export_file(tiny, expected_suffix=".pdf")

    assert missing_report["status"] == "失败"
    assert "输出文件不存在" in missing_report["errors"]
    assert tiny_report["status"] == "失败"
    assert any("输出文件过小" in error for error in tiny_report["errors"])



def test_docx_export_applies_reference_word_layout(tmp_path: Path) -> None:
    """DOCX 导出应接近用户参考 Word：A4、左右 3.17cm、标题居中、正文两字符缩进。"""
    output = tmp_path / "reference_style.docx"
    chapters = [
        {
            "number": 1,
            "title": "智慧工地建设总论",
            "content": "## 第一节 工程现场数字化转型背景\n### 一、行业发展与政策驱动\n#### （一）政策牵引\n正文段落用于验证导出版式。",
        }
    ]

    DocxExporter().export(chapters, str(output), title="智慧工地建设与工程现场数字化管理", author="测试作者", genre="专著")

    doc = Document(str(output))
    section = doc.sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7
    assert round(section.left_margin.cm, 2) == 3.17
    assert round(section.right_margin.cm, 2) == 3.17
    assert round(section.top_margin.cm, 2) == 2.54
    assert round(section.bottom_margin.cm, 2) == 2.54
    assert doc.styles["Heading 1"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert doc.styles["Heading 2"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert doc.styles["Normal"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert doc.styles["Heading 1"].font.size.pt == 20
    assert doc.styles["Heading 2"].font.size.pt == 15
    assert doc.styles["Normal"].font.size.pt == 10.5


def test_structured_json_formatted_word_export_builds_formal_sections(tmp_path: Path) -> None:
    """独立 JSON 脚本应生成封面、目录、前言、正文、结语、参考文献等正式 Word 结构。"""
    script_path = Path(__file__).resolve().parents[1] / "scripts" / "export_structured_json_to_docx.py"
    spec = importlib.util.spec_from_file_location("export_structured_json_to_docx", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    output = tmp_path / "formatted_word.docx"
    module.export_json_data_to_docx(
        {
            "title": "人工智能赋能高职图书馆阅读推广研究",
            "author": "测试作者",
            "preface": ["前言段落 AI 与阅读推广。"],
            "chapters": [
                {
                    "heading": "引言",
                    "sections": [
                        {
                            "heading": "研究背景",
                            "content": ["本节讨论 AI 技术与高职图书馆阅读推广的关系。"],
                            "level3": [
                                {
                                    "heading": "现实背景",
                                    "content": ["这是第一段。"],
                                    "level4": [
                                        {"heading": "政策牵引", "content": ["这是四级标题下的正文。"]}
                                    ],
                                }
                            ],
                        }
                    ],
                    "references": ["[1] 张三. 测试文献[M]. 北京: 测试出版社, 2024. https://example.com DOI:10.1234/test"],
                }
            ],
            "conclusion": ["结语段落。"],
            "global_references": ["# 参考文献", "[1] 李四. 全书参考文献[J]. 测试期刊, 2024. https://example.com/full"],
        },
        output,
    )

    assert output.exists()
    doc = Document(str(output))
    paragraph_texts = [p.text for p in doc.paragraphs]
    assert "目录" in paragraph_texts
    assert "前言" in paragraph_texts
    assert "结语" in paragraph_texts
    assert "参考文献" in paragraph_texts
    assert "Book Unnumbered Heading" in doc.styles
    assert "Reference" in doc.styles
    assert len(doc.sections) >= 4
    assert round(doc.sections[0].left_margin.cm, 2) == 3.18
    assert round(doc.sections[0].right_margin.cm, 2) == 3.18
    assert "Book Level4 Heading" in doc.styles
    assert "第一章　引言" in paragraph_texts
    assert "第一节　研究背景" in paragraph_texts
    assert "一、现实背景" in paragraph_texts
    assert "（一）政策牵引" in paragraph_texts
    assert not any("第1章" in text or "第1节" in text or "1.1" in text for text in paragraph_texts)
    assert not any("https://" in text or "DOI:" in text for text in paragraph_texts)

    heading_styles = {
        "Book Chapter Heading",
        "Book Section Heading",
        "Book Level3 Heading",
        "Book Level4 Heading",
        "Book Unnumbered Heading",
    }
    formatted_headings = [p for p in doc.paragraphs if p.style.name in heading_styles]
    assert formatted_headings
    for paragraph in formatted_headings:
        p_pr = paragraph._p.get_or_add_pPr()
        assert p_pr.find(module.qn("w:numPr")) is None
        assert p_pr.find(module.qn("w:pageBreakBefore")) is None
        assert p_pr.find(module.qn("w:keepNext")) is None
        assert p_pr.find(module.qn("w:keepLines")) is None

    assert doc.styles["Book Section Heading"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert doc.styles["Book Level3 Heading"].paragraph_format.first_line_indent.pt == 24
    assert doc.styles["Book Level4 Heading"].paragraph_format.first_line_indent.pt == 24
    assert doc.styles["Reference"].font.size.pt == 10.5
    assert doc.styles["Reference"].paragraph_format.left_indent.pt == 24
    assert doc.styles["Reference"].paragraph_format.first_line_indent.pt == -24



def test_academic_self_assessment_reports_percent_and_raw_score() -> None:
    content = "# 第一章 测试章节\n\n正文没有达到目标字数。"

    finalized, total, verdict = finalize_academic_chapter(content, target_words=3000)

    assert "/100" in verdict
    assert f"原始{total}/50" in verdict
    assert f"原始{total}/50" in finalized
    assert "总分：" in finalized


def test_export_readiness_report_flags_process_artifacts_and_short_chapters() -> None:
    chapters = [
        {
            "number": 1,
            "display_title": "第一章 测试章节",
            "content": "# 第一章 测试章节\n\n正文很短。\n\n【自评结果】\n总分：28/50 → 56/100",
        }
    ]

    report = QualityService.export_readiness_report(chapters)

    assert report["status"] == "不建议导出"
    assert report["readiness_score"] < 85
    assert report["duplicate_heading_count"] == 1
    assert report["process_artifact_count"] >= 1
    assert report["empty_chapter_count"] == 1
    assert any("过程" in item["问题"] or "评分" in item["问题"] for item in report["warnings"])


def test_export_readiness_report_accepts_clean_manuscript() -> None:
    clean_body = "\n\n".join(["　　这是符合专著语体的正文段落，围绕研究对象展开概念界定、机制分析和实践路径说明。" * 8 for _ in range(6)])
    chapters = [{"number": 1, "display_title": "第一章 测试章节", "content": clean_body}]

    report = QualityService.export_readiness_report(chapters)

    assert report["status"] == "通过"
    assert report["readiness_score"] == 100
    assert report["warnings"] == []


def test_manuscript_quality_report_flags_pseudo_citation_and_unsupported_claims() -> None:
    chapters = [
        {
            "number": 1,
            "display_title": "第一章 质量风险",
            "content": "## 第一节 背景\n\n　　已有研究表明智慧工地建设在2024年增长达到35%[3]。我们认为这是革命性的变化。",
        }
    ]

    report = QualityService.manuscript_quality_report(chapters, citations=[])

    assert report["status"] in {"建议复核", "不建议进入最终导出"}
    assert report["quality_score"] < 85
    assert report["pseudo_citation_flags"] >= 1
    assert report["unsupported_claim_flags"] >= 1
    assert report["style_flags"] >= 1
    assert any("伪引用" in item["问题"] or "无绑定引用" in item["问题"] for item in report["issues"])


def test_manuscript_quality_report_accepts_sufficient_academic_text_with_verified_citation() -> None:
    paragraph = "　　本节围绕智慧工地建设的概念边界、运行机制、实施路径和风险治理展开分析，强调工程现场数字化管理需要在组织体系、数据治理、技术集成和价值评估之间形成稳定协同。"
    content = "## 第一节 概念与机制\n\n" + "\n\n".join([paragraph for _ in range(10)])
    chapters = [{"number": 1, "display_title": "第一章 高质量正文", "content": content}]
    citations = [Citation(id="c-1", status="verified", formatted_ref="张三. 测试文献[M]. 2024.")]

    report = QualityService.manuscript_quality_report(chapters, citations=citations)

    assert report["status"] == "通过"
    assert report["quality_score"] >= 85
    assert report["issues"] == []


def test_source_citation_quality_pipeline_integration_chain() -> None:
    project = ProjectKnowledgeBase(
        project_name="commercial-demo",
        title="商业化专著",
        description="用于验证资料-证据-引用-质量-流水线链路。",
        category="专著",
        genre="测试",
        target_audience="研究者",
        book_length="短篇",
        tone="严谨",
        outline="# 第一章\n\n## 1.1 研究背景",
    )
    project.add_chapter(
        Chapter(
            chapter_number=1,
            title="第一章",
            actual_word_count=1500,
            status="completed",
            sections=[
                ChapterSection(
                    section_number="1.1",
                    title="研究背景",
                    summary="说明研究背景。",
                    level=2,
                    actual_word_count=800,
                    status="completed",
                    rag_query="研究背景",
                )
            ],
        )
    )
    project.add_source_document(
        SourceDocument(
            id="doc-1",
            title="测试资料",
            authors=["张三"],
            year="2024",
            source_type="M",
            doi="10.1234/demo",
            status="indexed",
            chunk_count=1,
        )
    )
    project.add_evidence_chunk(
        EvidenceChunk(
            id="chunk-1",
            document_id="doc-1",
            source="测试资料",
            text="研究背景表明，资料证据链能够提升专著写作中事实性论断的可追溯性。",
            page_start=12,
            chunk_index=1,
            score=0.92,
        )
    )

    source_service = SourceService()
    citation_service = CitationService()
    quality_service = QualityService()
    pipeline_service = PipelineService()

    source_summary = source_service.source_summary(project)
    document_coverage = source_service.document_coverage(project)
    citation = citation_service.bind_evidence_to_citation(
        project,
        "chunk-1",
        "资料证据链能够提升事实性论断的可追溯性。",
    )
    coverage_matrix = citation_service.coverage_matrix(project)
    quality = quality_service.summarize_project_quality(project)

    project.task_logs.append(TaskLog(task_type="export", status="completed", target="book.docx"))
    pipeline = pipeline_service.get_pipeline(project)
    pipeline_by_key = {stage["key"]: stage for stage in pipeline}

    assert source_summary["documents"] == 1
    assert source_summary["usable_chunks"] == 1
    assert document_coverage[0]["chunks"] == 1
    assert citation.status == "verified"
    assert coverage_matrix[0]["status"] == "verified"
    assert coverage_matrix[0]["has_evidence"] is True
    assert coverage_matrix[0]["has_source_document"] is True
    assert quality["citations"] >= 1
    assert "delivery" in pipeline_by_key
    assert pipeline_by_key["delivery"]["status"] == "in_progress"
    assert "引用核验" in pipeline_by_key["delivery"]["description"]
    assert pipeline_by_key["delivery"]["primary_route"] == "exports"
